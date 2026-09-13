from __future__ import annotations

import csv
import re
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "submission"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def main() -> None:
    pptx = SUBMISSION / "EvalPilot-Initial-Submission.pptx"
    pdf = SUBMISSION / "EvalPilot-Initial-Submission.pdf"
    form = SUBMISSION / "FORM_COPY.md"
    guide = SUBMISSION / "EVALUATOR_GUIDE.md"
    evidence = SUBMISSION / "EVIDENCE_INDEX.csv"
    review = SUBMISSION / "SIMULATED_JUDGE_REVIEW.md"

    for path in (pptx, pdf, form, guide, evidence, review):
        require(path.exists(), f"missing package file: {path.name}")

    prs = Presentation(pptx)
    require(len(prs.slides) == 9, f"PPTX must have 9 slides, found {len(prs.slides)}")
    require(all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() for slide in prs.slides),
            "every PPTX slide must have speaker notes")

    reader = PdfReader(pdf)
    require(len(reader.pages) == 9, f"PDF must have 9 pages, found {len(reader.pages)}")
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for token in ("EvalPilot", "8 / 26", "52 / 0", "4 / 4", "BLOCK", "$0.264697"):
        require(token in pdf_text, f"PDF is missing required token: {token}")

    form_text = form.read_text(encoding="utf-8")
    for token in (
        "逆熵智评 | NEGENTROPY LABS",
        "EvalPilot｜AI 应用回归评测与自主发布质量官",
        "https://www.bilibili.com/video/BV1mPYY6sE1k",
        "https://zqian6263-design.github.io/EvalPilot/",
    ):
        require(token in form_text, f"form copy is missing: {token}")

    with evidence.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    require(len(rows) >= 12, f"evidence index is too small: {len(rows)} rows")
    require({row["criterion"] for row in rows} == {"技术可行性", "市场可行性", "综合创新性", "AI 大模型结合", "赛道维度"},
            "evidence index must cover all five scoring criteria")

    for path in SUBMISSION.glob("*.md"):
        text = path.read_text(encoding="utf-8")
        require(not re.search(r"ghp_[A-Za-z0-9]+|sk-[A-Za-z0-9]+", text), f"possible secret in {path.name}")
        require("127.0.0.1" not in text, f"public form copy must not use localhost: {path.name}")

    print("Submission package verified: 5 criteria, 9 PPTX slides, 9 PDF pages.")


if __name__ == "__main__":
    main()